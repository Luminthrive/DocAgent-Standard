from typing import List

from pptx import Presentation
from langchain_core.documents import Document
from loguru import logger

from services.parsers.base_parser import BaseParser
from services.parsers import register_parser
from services.parsers.metadata_utils import build_base_metadata


@register_parser(".pptx")
class PPTXParser(BaseParser):
    """PPT 解析器 — 每张幻灯片一个 chunk"""

    async def parse(self, file_path: str) -> List[Document]:
        prs = Presentation(file_path)
        docs = []

        for idx, slide in enumerate(prs.slides):
            text_parts = []
            slide_title = ""
            has_notes = False

            # 提取幻灯片上的文字
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        para_text = para.text.strip()
                        if para_text:
                            text_parts.append(para_text)
                            if not slide_title:
                                slide_title = para_text

            # 提取备注
            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                notes_text = notes_frame.text.strip()
                if notes_text:
                    has_notes = True
                    text_parts.append(f"\n--- 备注 ---\n{notes_text}")

            page_content = "\n".join(text_parts)
            if not page_content.strip():
                continue

            meta = build_base_metadata(
                file_path, ".pptx", chunk_index=len(docs), total_chunks=0,
                doc_title=slide_title,
                location_ref=f"幻灯片第{idx + 1}页",
                extra={
                    "slide_number": idx + 1,
                    "slide_title": slide_title,
                    "has_notes": has_notes,
                    "_parsed_raw": page_content,
                },
            )
            docs.append(Document(page_content=page_content, metadata=meta))

        # 回填 total_chunks
        for c in docs:
            c.metadata["total_chunks"] = len(docs)

        logger.info(f"PPTX 解析完成: {len(docs)} slides")
        return docs
